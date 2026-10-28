import json, re
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
c=json.loads(Path('tmp/report-electronics/content.json').read_text(encoding='utf-8'))
d=Document()
for sec in d.sections:
 sec.page_width=Cm(21); sec.page_height=Cm(29.7)
 sec.top_margin=Cm(2.6); sec.bottom_margin=Cm(2.6); sec.left_margin=Cm(3); sec.right_margin=Cm(3)
 sec.header_distance=Cm(1.2); sec.footer_distance=Cm(1.3)
for name in ['Normal','Title','Heading 1','Footer']:
 s=d.styles[name]; s.font.name='Times New Roman'; s.font.size=Pt(12); s.font.color.rgb=RGBColor(0,0,0)
 s.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体')
 pf=s.paragraph_format; pf.line_spacing=1.2; pf.space_before=Pt(0); pf.space_after=Pt(6)
 pf.widow_control=True
s=d.styles['Heading 1']; s.font.bold=True; s.paragraph_format.space_before=Pt(12); s.paragraph_format.space_after=Pt(6)
s.paragraph_format.keep_with_next=True
# Cover
p=d.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before=Pt(125);p.paragraph_format.space_after=Pt(30)
r=p.add_run('科技论文阅读报告');r.bold=True
p=d.add_paragraph(c['title']);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(90)
for text in ['姓名  '+c['name'], '学号  '+c['id']]:
 p=d.add_paragraph(text);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(20)
sec=d.add_section(WD_SECTION_START.NEW_PAGE)
sec.footer.is_linked_to_previous=False
num=OxmlElement('w:pgNumType');num.set(qn('w:start'),'1');sec._sectPr.append(num)
p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run();f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');r._r.addnext(f)
def body(text,indent=True):
 p=d.add_paragraph(text)
 if indent:p.paragraph_format.first_line_indent=Pt(24)
 return p
body(c['intro'])
for i,s in enumerate(c['sections']):
 p=d.add_paragraph(s['h'],style='Heading 1')
 if i in [3,6]:p.paragraph_format.page_break_before=True
 for text in s['p']:body(text)
p=d.add_paragraph('阅读文献',style='Heading 1')
body(c['reference'],False)
# Explicit fonts and size on all runs, including Latin metadata.
for p in d.paragraphs:
 for r in p.runs:
  r.font.name='Times New Roman';r.font.size=Pt(12)
  r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体')
out=Path('output/reports/多模态情感识别论文阅读报告.docx')
# Remove inherited template decoration from all paragraph styles.
for element in list(d.styles.element.iter(qn('w:pBdr'))):
 element.getparent().remove(element)
for element in list(d.element.iter(qn('w:pBdr'))):
 element.getparent().remove(element)
d.save(out)
print(out)
print('Body Han characters:',len(re.findall(r'[\u3400-\u9fff]',c['intro']+''.join(s['h']+''.join(s['p']) for s in c['sections']))))
